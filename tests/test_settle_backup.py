"""v1.3.0：结算闭环与备份恢复。"""
from __future__ import annotations

from pathlib import Path

from core.models import KIND_GROUPBUY, Item
from core.repository import items as item_repo
from core.repository import signups as signup_repo


def _new_item(title: str = "结算测试拼单") -> Item:
    it = Item(title=title, kind=KIND_GROUPBUY, price=10.0, unit="份",
              quota=50, source_hash=f"hash-{title}")
    it.id = item_repo.create(it)   # create 只返回 id，不回填对象
    return it


# ===========================================================================
# 结算
# ===========================================================================

def test_settlement_summary_splits_by_settled():
    """结清一部分人后，人数与份数都要按结清/待结清拆对。"""
    it = _new_item()
    signup_repo.batch_add(it.id, [("甲", 2.0), ("乙", 1.0), ("丙", 3.0)])

    before = signup_repo.settlement_summary(it.id)
    assert before["people"] == 3
    assert before["done_people"] == 0 and before["wait_people"] == 3

    rows = signup_repo.list_for(it.id)
    signup_repo.set_settled(it.id, True, [rows[0].id])

    after = signup_repo.settlement_summary(it.id)
    assert after["done_people"] == 1 and after["wait_people"] == 2
    assert after["done_qty"] == 2.0 and after["wait_qty"] == 4.0
    # 拆分后总量不能变
    assert abs((after["done_qty"] + after["wait_qty"]) - before["qty"]) < 1e-6


def test_settled_can_be_flipped_back():
    """误标了要能取消——批量与整条两个口径都要生效。"""
    it = _new_item("结算反复拼单")
    signup_repo.batch_add(it.id, [("甲", 1.0), ("乙", 1.0)])

    signup_repo.set_settled(it.id, True)          # 整条全标
    assert signup_repo.settlement_summary(it.id)["wait_people"] == 0

    signup_repo.set_settled(it.id, False)         # 整条全取消
    assert signup_repo.settlement_summary(it.id)["done_people"] == 0


def test_export_pending_only_contains_unsettled():
    """催款清单里不能混进已经结清的人。"""
    from openpyxl import load_workbook

    from exports import excel

    it = _new_item("结算导出拼单")
    signup_repo.batch_add(it.id, [("甲", 1.0), ("乙", 1.0), ("丙", 1.0)])
    rows = signup_repo.list_for(it.id)
    signup_repo.set_settled(it.id, True, [rows[0].id])

    path = excel.export_settlement(it, signup_repo.list_for(it.id),
                                   only_unsettled=True)
    assert Path(path).exists()

    ws = load_workbook(path).active
    names = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(names) == 2, f"待结清应有 2 人，实际 {len(names)}：{names}"


# ===========================================================================
# 备份与恢复
# ===========================================================================

def test_backup_roundtrip_restores_data():
    """备份 → 删一条 → 恢复 → 条目数与报名都得原样回来。"""
    from core import backup

    it = _new_item("备份往返拼单")
    signup_repo.batch_add(it.id, [("甲", 5.0)])

    before = len(item_repo.list_items(include_archived=True))
    zip_path = backup.create_backup()
    assert Path(zip_path).exists()

    item_repo.delete(it.id)
    assert len(item_repo.list_items(include_archived=True)) == before - 1

    ok, msg = backup.restore_backup(zip_path)
    assert ok, msg
    assert len(item_repo.list_items(include_archived=True)) == before

    restored = [x for x in item_repo.list_items(include_archived=True)
                if x.source_hash == "hash-备份往返拼单"]
    assert restored, "恢复后找不到被删的那条"
    assert signup_repo.total_qty(restored[0].id) == 5.0, "报名记录没跟着回来"


def test_backup_manifest_matches_db():
    """备份清单里写的条目数要和库里一致，恢复前才能给用户看清楚。"""
    from core import backup

    _new_item("清单核对拼单")
    zip_path = backup.create_backup()
    manifest = backup.read_manifest(zip_path)
    assert manifest, "备份包里没有清单"
    assert manifest["counts"]["items"] == len(item_repo.list_items(include_archived=True))
    assert manifest["app_version"]


def test_restore_rejects_invalid_file():
    """拿错文件不能把现有数据搞坏。"""
    import tempfile

    from core import backup

    junk = Path(tempfile.mkdtemp()) / "not-a-backup.zip"
    junk.write_text("这不是备份", encoding="utf-8")

    ok, msg = backup.restore_backup(junk)
    assert not ok and msg, "无效备份必须被拒绝且给出原因"
