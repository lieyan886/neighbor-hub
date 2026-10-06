"""v1.7.0：价格历史 / 接龙里的「我」记给团长 / 提货清单补备注列。"""
from __future__ import annotations

import pytest

from core import config
from core.models import Item, Signup, STATUS_ACTIVE
from core.repository import items as item_repo
from core.repository import prices, signups as signup_repo
from exports import excel
from render import copywriter


def _mk(title: str, price: float | None = None, **kw) -> Item:
    kw.setdefault("status", STATUS_ACTIVE)
    kw.setdefault("kind", "deal")
    it = Item(title=title, unit="份", price=price,
              source_hash=f"v17-{title}", **kw)
    it.id = item_repo.create(it)
    return it


# ===========================================================================
# 价格历史
# ===========================================================================

def test_price_recorded_on_create(solo_data):
    """入库时的价格就是历史第一笔。"""
    it = _mk("车厘子 2斤", price=168.0)
    seq = prices.series(it.id)
    assert len(seq) == 1, f"入库该记一笔，实际 {seq}"
    assert seq[0][1] == 168.0


def test_price_recorded_when_changed(solo_data):
    """改价要留痕，团长靠它判断真降还是先涨后降。"""
    it = _mk("牛肉卷", price=168.0)
    it.price = 158.0
    item_repo.update(it)

    seq = prices.series(it.id)
    assert len(seq) == 2
    assert seq[-1][1] == 158.0
    s = prices.summary(it.id)
    assert s["high"] == 168.0 and s["low"] == 158.0 and s["now"] == 158.0


def test_price_not_recorded_when_unchanged(solo_data):
    """反复保存但价格没变，不该刷出一堆重复点。"""
    it = _mk("牛奶", price=59.0)
    item_repo.update(it)
    item_repo.update(it)
    assert len(prices.series(it.id)) == 1


def test_no_price_no_history(solo_data):
    """没填价格的条目不进价格历史（画不出点，还会污染下拉）。"""
    it = _mk("纯福利", price=None)
    assert prices.series(it.id) == []


def test_price_subjects_listed(solo_data):
    """看板下拉要能列出「哪些内容攒下了价格历史」。"""
    it = _mk("橙子", price=39.0)
    subs = prices.subjects()
    assert any(k == f"item:{it.id}" for k, _, _ in subs)


def test_price_history_survives_rename(solo_data):
    """改标题不该把价格历史丢掉（按 item_id 关联，不按标题）。"""
    it = _mk("原名", price=88.0)
    it.title = "改名了"
    item_repo.update(it)
    assert len(prices.series(it.id)) == 1


# ===========================================================================
# 接龙里的「我」
# ===========================================================================

def test_solitaire_self_name_replaces_pronoun():
    """「我要2份」落到团长名下，而不是被当成认不出来。"""
    rows = copywriter.parse_solitaire("我要2份\n李四 3份", self_name="3栋小李")
    assert ("3栋小李", 2.0) in rows
    assert ("李四", 3.0) in rows
    assert "我" not in [n for n, _ in rows]


def test_solitaire_without_self_name_still_skips():
    """没给团长名，「我」照样进「认不出来」，不静默记一条叫「我」的。"""
    res = copywriter.parse_solitaire_detail("我要2份")
    assert res.rows == []
    assert res.skipped


def test_solitaire_self_cancel():
    """「我不要了」也要认成团长自己取消。"""
    res = copywriter.parse_solitaire_detail("我不要了", self_name="团长")
    assert ("团长", None) in res.adjustments


def test_solitaire_normal_names_unaffected():
    """正常人名不受 self_name 影响。"""
    rows = copywriter.parse_solitaire("1. 张三 2份\n王五 1", self_name="团长")
    assert set(n for n, _ in rows) == {"张三", "王五"}


# ===========================================================================
# 提货清单备注列
# ===========================================================================

def test_packing_sheet_has_note_column(solo_data):
    """志愿者手里是打印出来的纸，备注必须印在上面。"""
    it = _mk("草莓", price=45.0, kind="groupbuy")
    signup_repo.add(Signup(item_id=it.id, name="甲", qty=2, unit="盒",
                           note="要红颜不要奶油"))
    path = excel.export_groupbuy_packing([it], {it.id: signup_repo.list_for(it.id)})
    from openpyxl import load_workbook

    ws = load_workbook(path).active
    header = [c.value for c in ws[1]]
    assert "备注" in header, f"表头缺少备注列：{header}"
    idx = header.index("备注")
    notes = [ws.cell(row=2, column=idx + 1).value]
    assert notes == ["要红颜不要奶油"]
