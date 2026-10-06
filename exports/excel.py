"""Excel 导出：报名表、内容总表、拼单结算清单（只列数量不含金额逻辑）。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core import config, privacy
from core.models import Item, KIND_LABELS, STATUS_LABELS, kind_label, status_label
from core.utils import format_price

_HEADER_FILL = PatternFill("solid", fgColor="2F4F6F")
_HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
_CELL_FONT = Font(name="微软雅黑", size=10)
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(vertical="center", wrap_text=True)
_CENTER = Alignment(horizontal="center", vertical="center")


def _style_header(ws, widths: dict[str, int]) -> None:
    for idx, width in enumerate(widths.values(), start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 26
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _BORDER
    ws.freeze_panes = "A2"


def _append_rows(ws, rows: list[list], wrap_cols: set[int] = frozenset()) -> None:
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = _CELL_FONT
            cell.border = _BORDER
            cell.alignment = _WRAP if cell.column in wrap_cols else Alignment(vertical="center")


def export_signups(item: Item, rows: list, output_dir: str | Path | None = None) -> str:
    """导出某条内容的报名/接龙明细。"""
    config.ensure_dirs()
    wb = Workbook()
    ws = wb.active
    ws.title = "报名明细"
    title = f"{kind_label(item.kind)}｜{item.title}"
    ws.append(["序号", "姓名/昵称", "联系方式", "数量", "单位", "备注", "已结清", "登记时间"])

    # 导出文件可能被转发/留档，一律按 export 口径脱敏
    rows = privacy.scrub_rows(rows, scope="export")

    for idx, s in enumerate(rows, 1):
        ws.append([idx, s.name, s.contact, s.qty, s.unit, s.note,
                   "是" if s.settled else "否", s.created_at])
    total = sum(float(s.qty or 0) for s in rows)
    ws.append(["", "合计", "", total, item.unit or "份", "", "", ""])

    _style_header(ws, {"序号": 6, "姓名/昵称": 16, "联系方式": 16, "数量": 8,
                       "单位": 8, "备注": 24, "已结清": 8, "登记时间": 18})
    last = ws.max_row
    for cell in ws[last]:
        cell.font = Font(name="微软雅黑", size=10, bold=True)

    dest = Path(output_dir or config.EXPORT_DIR)
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / f"报名明细_{_safe(item.title)}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return str(path)


def export_items(items: list[Item], output_dir: str | Path | None = None) -> str:
    """导出内容总表，便于线下贴 pal 或做月度复盘。"""
    config.ensure_dirs()
    wb = Workbook()
    ws = wb.active
    ws.title = "内容总表"
    ws.append(["ID", "类型", "标题", "商户/主办", "地点", "单价", "原价", "单位",
               "目标量", "活动时间", "截止时间", "状态", "标签", "链接"])

    for it in items:
        ws.append([
            it.id, kind_label(it.kind), it.title, it.merchant, it.location,
            it.price if it.price is not None else "",
            it.origin_price if it.origin_price is not None else "",
            it.unit, it.quota or "", it.event_at, it.deadline,
            status_label(it.status), it.tags, it.url,
        ])

    _style_header(ws, {"ID": 6, "类型": 10, "标题": 36, "商户/主办": 16, "地点": 16,
                       "单价": 8, "原价": 8, "单位": 6, "目标量": 8, "活动时间": 16,
                       "截止时间": 16, "状态": 10, "标签": 18, "链接": 30})

    dest = Path(output_dir or config.EXPORT_DIR)
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / f"内容总表_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return str(path)


def export_groupbuy_packing(items: list[Item], signups_map: dict[int, list],
                            output_dir: str | Path | None = None) -> str:
    """拼单提货清单：按人汇总，方便到货时核对（不涉及金额）。"""
    config.ensure_dirs()
    wb = Workbook()
    ws = wb.active
    ws.title = "提货清单"
    # v1.7.0：加「备注」列 —— 自提时最常问的就是「他要的是辣口还是原味」，
    # 以前这列只有界面上看得到，打印出来的清单上没有，志愿者还得回头翻软件
    ws.append(["所属拼单", "姓名/昵称", "联系方式", "数量", "单位", "备注", "已提货"])

    for it in items:
        # 提货清单会给到团长/志愿者，同样按 export 口径脱敏
        for s in privacy.scrub_rows(signups_map.get(it.id or -1, []), scope="export"):
            ws.append([it.title, s.name, s.contact, s.qty,
                       s.unit or it.unit or "份", s.note, ""])

    _style_header(ws, {"所属拼单": 30, "姓名/昵称": 16, "联系方式": 16,
                       "数量": 8, "单位": 8, "备注": 24, "已提货": 10})

    dest = Path(output_dir or config.EXPORT_DIR)
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / f"提货清单_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return str(path)


def export_settlement(item: Item, rows: list, only_unsettled: bool = False,
                      output_dir: str | Path | None = None) -> str:
    """结算清单：谁结清了、谁还没。

    工具不碰钱，这里只汇总「份数」与「结清状态」，金额由团长按群收款自己核。
    only_unsettled=True 时只导出待结清的人，用于催一轮。
    """
    config.ensure_dirs()
    from core.repository import signups as signup_repo

    summary = signup_repo.settlement_summary(item.id or 0)

    wb = Workbook()
    ws = wb.active
    ws.title = "待结清" if only_unsettled else "结算清单"
    ws.append(["序号", "姓名/昵称", "联系方式", "数量", "单位", "已结清", "备注", "登记时间"])

    # 清单可能被转发给志愿者核对，一律按 export 口径脱敏
    picked = [s for s in privacy.scrub_rows(rows, scope="export")
              if (not s.settled) or not only_unsettled]

    for idx, s in enumerate(picked, 1):
        ws.append([idx, s.name, s.contact, s.qty, s.unit or item.unit or "份",
                   "是" if s.settled else "否", s.note, s.created_at])

    unit = item.unit or "份"
    ws.append(["", "合计", "", sum(float(s.qty or 0) for s in picked), unit, "", "", ""])
    ws.append(["", f"共 {summary['people']} 人 {summary['qty']:g}{unit}"
                   f"（已结清 {summary['done_people']} 人 / 待结清 {summary['wait_people']} 人）",
               "", "", "", "", "", ""])

    _style_header(ws, {"序号": 6, "姓名/昵称": 16, "联系方式": 16, "数量": 8,
                       "单位": 8, "已结清": 8, "备注": 24, "登记时间": 18})
    for row in ws.iter_rows(min_row=ws.max_row - 1):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10, bold=True)

    dest = Path(output_dir or config.EXPORT_DIR)
    dest.mkdir(parents=True, exist_ok=True)
    kind = "待结清清单" if only_unsettled else "结算清单"
    path = dest / f"{kind}_{_safe(item.title)}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    wb.save(path)
    return str(path)


def _safe(text: str) -> str:
    return "".join(c for c in (text or "")[:20] if c not in '\\/:*?"<>|') or "未命名"
